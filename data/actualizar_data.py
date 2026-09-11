"""
Convierte los Excel de data/ en data/data.js para que el tablero los lea.

El tablero se abre como archivo o desde GitHub Pages, y el navegador no puede leer un Excel,
así que cada vez que cambie un Excel hay que correr, desde la carpeta dashboard:

    python data/actualizar_data.py

Flujo completo:
    python data/scrapear_mef.py      (descarga del MEF y genera data.xlsx: ranking de Sector)
    python data/actualizar_data.py   (lee data.xlsx y "Datos para PPT.xlsx" y genera data.js)
    git add, commit y push

data.xlsx
  Ranking            Sector | PIM | Devengado | Avance %          (obligatoria, total nacional)
  Ranking regiones   Código | Departamento | Sector | PIM | Devengado | Avance %   (opcional)
  Info               Dato | Valor                                  (opcional)

Datos para PPT.xlsx (opcional; si no existe, esos gráficos usan los datos de ejemplo)
  Obras en ejecución              SUBSECTOR | EJECUCION | DEPARTAMENTO | OBRAS | COSTO DE INVERSIÓN |
                                  POBLACIÓN BENEFICIARIA | CONEXIONES NUEVAS DE AGUA | CONEXIONES NUEVAS DE ALCANTARILLADO/UBS
  Inversiones en vivienda         Departamentos | Bonos desembolsados | Bono monto | Bono población |
                                  Vivienda rural construida | Vivienda rural población | Vivienda rural monto | Predios registrados
  Expedientes técnicos en PRESET  REGIÓN | TOTAL | POBLACION | INVERSION | APTO | EN EVALUACION  (por región)
  Obras paralizadas / Obras reactivadas  REGIÓN | TOTAL | ... | DIRECTA | ... | TRANSFERENCIA | ... | SANEAMIENTO | Población | Monto |
                                  VIVIENDA | Población | Monto | CONEXIONES NUEVAS DE AGUA  (directa/transferencia solo en total)
  Cierre de brecha = conexiones nuevas (ejecución: agua + alcantarillado; paralizadas y reactivadas: agua).
  Nacional = suma de los departamentos. Montos en soles; en data.js van en millones.
"""
import datetime
import io
import json
import os
import re
import sys
import unicodedata

try:
    import openpyxl
except ImportError:
    sys.exit('Falta openpyxl. Instálalo con: pip install openpyxl')

AQUI = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(AQUI, 'data.xlsx')
PPT = os.path.join(AQUI, 'Datos para PPT.xlsx')
SALIDA = os.path.join(AQUI, 'data.js')

DEPARTAMENTOS = ['Amazonas', 'Áncash', 'Apurímac', 'Arequipa', 'Ayacucho', 'Cajamarca', 'Callao', 'Cusco', 'Huancavelica',
                 'Huánuco', 'Ica', 'Junín', 'La Libertad', 'Lambayeque', 'Lima', 'Loreto', 'Madre de Dios', 'Moquegua',
                 'Pasco', 'Piura', 'Puno', 'San Martín', 'Tacna', 'Tumbes', 'Ucayali']


def limpio(v):
    return str(v).replace('\xa0', ' ').strip() if v is not None else ''


def clave(v):
    """Texto comparable: sin tildes, en mayúsculas y con espacios simples ("LIMA ", "Huánuco" → "LIMA", "HUANUCO")."""
    t = unicodedata.normalize('NFD', limpio(v))
    t = ''.join(c for c in t if unicodedata.category(c) != 'Mn')
    return re.sub(r'\s+', ' ', t).strip().upper()


MAPA = {clave(n): n for n in DEPARTAMENTOS}
MAPA['PROVINCIA CONSTITUCIONAL DEL CALLAO'] = 'Callao'


def numero(v):
    if v is None or v == '':
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = limpio(v).replace(',', '').replace('%', '')
    return float(t) if re.fullmatch(r'-?\d+(\.\d+)?', t) else None


def num0(v):
    n = numero(v)
    return n if n is not None else 0.0


def cantidad_texto(v):
    """'5 mil millones' → 5e9 · '4 millones' → 4e6 · 1234 → 1234."""
    if isinstance(v, (int, float)):
        return float(v)
    t = clave(v).replace(',', '.')
    m = re.search(r'\d+(\.\d+)?', t)
    if not m:
        return 0.0
    base = float(m.group(0))
    if 'MIL MILLON' in t:
        return base * 1e9
    if 'MILLON' in t:
        return base * 1e6
    if 'MIL' in t:
        return base * 1e3
    return base


# ---------------------------------------------------------------- data.xlsx (Sector)
def columnas(ws, nombres):
    cab = [limpio(c).lower() for c in next(ws.iter_rows(values_only=True))]
    idx = {}
    for n in nombres:
        pos = [i for i, c in enumerate(cab) if c.startswith(n)]
        if not pos:
            sys.exit('Falta la columna "' + n + '" en la hoja ' + ws.title)
        idx[n] = pos[0]
    return idx


def fila_sector(f, i):
    texto = limpio(f[i['sector']])
    codigo, _, nombre = texto.partition(':')
    if not nombre:
        codigo, nombre = '', texto
    pim, dev, av = numero(f[i['pim']]), numero(f[i['devengado']]), numero(f[i['avance']])
    if av is None and pim:
        av = (dev or 0) / pim
    if av is not None and av <= 1.5:  # viene como fracción
        av = av * 100
    return {'codigo': codigo.strip(), 'sector': nombre.strip(), 'pim': round(pim or 0), 'devengado': round(dev or 0),
            'avance': round(av, 1) if av is not None else None}


def leer_ranking(wb):
    if 'Ranking' not in wb.sheetnames:
        sys.exit('No existe la hoja "Ranking" en data.xlsx')
    ws = wb['Ranking']
    i = columnas(ws, ['sector', 'pim', 'devengado', 'avance'])
    return [fila_sector(f, i) for f in list(ws.iter_rows(values_only=True))[1:] if f and limpio(f[i['sector']])]


def leer_regiones(wb):
    if 'Ranking regiones' not in wb.sheetnames:
        return {}
    ws = wb['Ranking regiones']
    i = columnas(ws, ['departamento', 'sector', 'pim', 'devengado', 'avance'])
    por_region = {}
    for f in list(ws.iter_rows(values_only=True))[1:]:
        if not f or not limpio(f[i['sector']]) or not limpio(f[i['departamento']]):
            continue
        por_region.setdefault(limpio(f[i['departamento']]), []).append(fila_sector(f, i))
    return por_region


def leer_info(wb):
    if 'Info' not in wb.sheetnames:
        return {}
    return {limpio(f[0]).lower(): limpio(f[1]) for f in list(wb['Info'].iter_rows(values_only=True))[1:] if f and f[0]}


# ---------------------------------------------------------------- Datos para PPT.xlsx
def hoja(wb, inicio):
    for n in wb.sheetnames:
        if clave(n).startswith(clave(inicio)):
            return wb[n]
    return None


def indice(cab, *palabras):
    for i, c in enumerate(cab):
        if all(p in c for p in palabras):
            return i
    sys.exit('Falta una columna con: ' + ' '.join(palabras))


def departamento(v, avisos):
    dep = MAPA.get(clave(v))
    if not dep and clave(v):
        avisos.add(limpio(v))
    return dep


def leer_obras(ws, avisos):
    filas = [f for f in ws.iter_rows(values_only=True) if f and any(x not in (None, '') for x in f)]
    cab = [clave(c) for c in filas[0]]
    i_sub, i_eje, i_dep = indice(cab, 'SUBSECTOR'), indice(cab, 'EJECUCION'), indice(cab, 'DEPARTAMENTO')
    i_obr, i_cos, i_pob = indice(cab, 'OBRAS'), indice(cab, 'COSTO'), indice(cab, 'POBLACION')
    i_agu, i_alc = indice(cab, 'CONEXIONES', 'AGUA'), indice(cab, 'CONEXIONES', 'ALCANTARILLADO')

    def vacio():
        return {'total': 0, 'directa': 0, 'transferencia': 0, 'monto': 0.0, 'poblacion': 0,
                'conexionesAgua': 0, 'conexionesAlcantarillado': 0}
    reg = {n: {'saneamiento': vacio(), 'vivienda': vacio()} for n in ['Nacional'] + DEPARTAMENTOS}
    for f in filas[1:]:
        dep = departamento(f[i_dep], avisos)
        if not dep:
            continue
        sub = 'vivienda' if clave(f[i_sub]).startswith('VIVIENDA') else 'saneamiento'
        eje = 'directa' if clave(f[i_eje]).startswith('DIRECTA') else 'transferencia'
        for destino in (dep, 'Nacional'):
            s = reg[destino][sub]
            n = int(num0(f[i_obr]))
            s['total'] += n
            s[eje] += n
            s['monto'] += num0(f[i_cos]) / 1e6
            s['poblacion'] += int(num0(f[i_pob]))
            s['conexionesAgua'] += int(num0(f[i_agu]))
            s['conexionesAlcantarillado'] += int(num0(f[i_alc]))
    for r in reg.values():
        for s in r.values():
            s['monto'] = round(s['monto'], 1)
            s['conexiones'] = s['conexionesAgua'] + s['conexionesAlcantarillado']  # cierre de brecha = agua + alcantarillado
    return reg


def leer_vivienda(ws, avisos):
    filas = [f for f in ws.iter_rows(values_only=True) if f and any(x not in (None, '') for x in f)]
    cab = [clave(c) for c in filas[0]]
    i_dep = indice(cab, 'DEPARTAMENTO')
    i_bc, i_bm, i_bp = indice(cab, 'BONOS', 'DESEMBOLSADOS'), indice(cab, 'BONO', 'MONTO'), indice(cab, 'BONO', 'POBLACION')
    i_rc, i_rp, i_rm = indice(cab, 'RURAL', 'CONSTRUIDA'), indice(cab, 'RURAL', 'POBLACION'), indice(cab, 'RURAL', 'MONTO')
    i_pr = indice(cab, 'PREDIOS')

    def vacio():
        return {'bonos': {'cantidad': 0, 'monto': 0.0, 'poblacion': 0},
                'viviendaRural': {'cantidad': 0, 'monto': 0.0, 'poblacion': 0},
                'predios': {'cantidad': 0}}
    reg = {n: vacio() for n in ['Nacional'] + DEPARTAMENTOS}
    for f in filas[1:]:
        dep = departamento(f[i_dep], avisos)
        if not dep:
            continue
        for destino in (dep, 'Nacional'):
            v = reg[destino]
            v['bonos']['cantidad'] += int(num0(f[i_bc]))
            v['bonos']['monto'] += num0(f[i_bm]) / 1e6
            v['bonos']['poblacion'] += int(num0(f[i_bp]))
            v['viviendaRural']['cantidad'] += int(num0(f[i_rc]))
            v['viviendaRural']['monto'] += num0(f[i_rm]) / 1e6
            v['viviendaRural']['poblacion'] += int(num0(f[i_rp]))
            v['predios']['cantidad'] += int(num0(f[i_pr]))
    for v in reg.values():
        v['bonos']['monto'] = round(v['bonos']['monto'], 1)
        v['viviendaRural']['monto'] = round(v['viviendaRural']['monto'], 1)
    return reg


def leer_preset(ws, avisos):
    """Tabla por región: REGIÓN | TOTAL | POBLACION | INVERSION | APTO | EN EVALUACION (acepta la errata "EVALAUCION").
    Nacional = suma de los departamentos; si hay fila "Total general" se compara y se avisa si no coincide.
    Si la hoja viene en el formato antiguo (Dato | Valor, solo total nacional), devuelve solo Nacional."""
    filas = [f for f in ws.iter_rows(values_only=True) if f and any(x not in (None, '') for x in f)]
    cab = [clave(c) for c in filas[0]]
    if not any(c.startswith('REGION') or c.startswith('DEPARTAMENTO') for c in cab):
        kv = {clave(f[0]): f[1] for f in filas if f[0] not in (None, '')}
        def valor(prefijo):
            return next((v for k, v in kv.items() if k.startswith(prefijo)), None)
        return {'Nacional': {'total': int(num0(valor('TOTAL'))), 'evaluacion': int(num0(valor('EN EVALUACION'))), 'aptos': int(num0(valor('APTOS'))),
                             'monto': round(cantidad_texto(valor('MONTO')) / 1e6, 1), 'poblacion': int(cantidad_texto(valor('POBLACION')))}}
    i_reg = next(i for i, c in enumerate(cab) if c.startswith('REGION') or c.startswith('DEPARTAMENTO'))
    i_tot, i_pob, i_inv, i_apt = indice(cab, 'TOTAL'), indice(cab, 'POBLACION'), indice(cab, 'INVERSION'), indice(cab, 'APTO')
    i_eva = next((i for i, c in enumerate(cab) if c.startswith('EN EVAL')), None)
    if i_eva is None:
        sys.exit('Falta la columna EN EVALUACIÓN en la hoja de PRESET')

    def vacio():
        return {'total': 0, 'evaluacion': 0, 'aptos': 0, 'monto': 0.0, 'poblacion': 0}
    reg = {n: vacio() for n in ['Nacional'] + DEPARTAMENTOS}
    fila_total = None
    for f in filas[1:]:
        if clave(f[i_reg]).startswith('TOTAL'):
            fila_total = f
            continue
        dep = departamento(f[i_reg], avisos)
        if not dep:
            continue
        for destino in (dep, 'Nacional'):
            q = reg[destino]
            q['total'] += int(num0(f[i_tot]))
            q['evaluacion'] += int(num0(f[i_eva]))
            q['aptos'] += int(num0(f[i_apt]))
            q['monto'] += num0(f[i_inv]) / 1e6
            q['poblacion'] += int(num0(f[i_pob]))
    for q in reg.values():
        q['monto'] = round(q['monto'], 1)
    if fila_total is not None:
        n = reg['Nacional']
        esperado = (int(num0(fila_total[i_tot])), int(num0(fila_total[i_eva])), int(num0(fila_total[i_apt])), int(num0(fila_total[i_pob])))
        obtenido = (n['total'], n['evaluacion'], n['aptos'], n['poblacion'])
        if esperado != obtenido:
            print('Aviso: la fila "Total general" de PRESET no coincide con la suma de regiones:', esperado, 'frente a', obtenido)
        else:
            print('Control PRESET: la suma de regiones coincide con "Total general".')
    return reg


def leer_modalidad(ws, avisos):
    """Obras paralizadas / reactivadas por región: DIRECTA y TRANSFERENCIA (solo en total), SANEAMIENTO y VIVIENDA
    (cantidad, población y monto) y CONEXIONES NUEVAS DE AGUA (y de alcantarillado, si la hoja la trae).
    El Excel no desglosa directa/transferencia por sector: van como total. Las conexiones van a saneamiento."""
    filas = [f for f in ws.iter_rows(values_only=True) if f and any(x not in (None, '') for x in f)]
    cab = [clave(c) for c in filas[0]]

    def exacto(nombre):
        for i, c in enumerate(cab):
            if c == nombre:
                return i
        sys.exit('Falta la columna ' + nombre + ' en la hoja ' + ws.title)
    i_reg = next(i for i, c in enumerate(cab) if c.startswith('REGION') or c.startswith('DEPARTAMENTO'))
    i_dir, i_tra = exacto('DIRECTA'), exacto('TRANSFERENCIA')
    i_sc, i_sp, i_sm = exacto('SANEAMIENTO'), indice(cab, 'POBLACION', 'SANEAMIENTO'), indice(cab, 'MONTO', 'SANEAMIENTO')
    i_vc, i_vp, i_vm = exacto('VIVIENDA'), indice(cab, 'POBLACION', 'VIVIENDA'), indice(cab, 'MONTO', 'VIVIENDA')
    i_agu = indice(cab, 'CONEXIONES', 'AGUA')
    i_alc = next((i for i, c in enumerate(cab) if 'CONEXIONES' in c and 'ALCANTARILLADO' in c), None)

    def sector():
        return {'total': 0, 'directa': None, 'transferencia': None, 'monto': 0.0, 'poblacion': 0, 'conexiones': 0}
    reg = {n: {'saneamiento': sector(), 'vivienda': sector(), 'directaTotal': 0, 'transferenciaTotal': 0} for n in ['Nacional'] + DEPARTAMENTOS}
    for f in filas[1:]:
        if clave(f[i_reg]).startswith('TOTAL'):
            continue
        dep = departamento(f[i_reg], avisos)
        if not dep:
            continue
        con = int(num0(f[i_agu])) + (int(num0(f[i_alc])) if i_alc is not None else 0)
        for destino in (dep, 'Nacional'):
            r = reg[destino]
            r['directaTotal'] += int(num0(f[i_dir]))
            r['transferenciaTotal'] += int(num0(f[i_tra]))
            r['saneamiento']['total'] += int(num0(f[i_sc]))
            r['saneamiento']['poblacion'] += int(num0(f[i_sp]))
            r['saneamiento']['monto'] += num0(f[i_sm]) / 1e6
            r['saneamiento']['conexiones'] += con
            r['vivienda']['total'] += int(num0(f[i_vc]))
            r['vivienda']['poblacion'] += int(num0(f[i_vp]))
            r['vivienda']['monto'] += num0(f[i_vm]) / 1e6
    for r in reg.values():
        for k in ('saneamiento', 'vivienda'):
            r[k]['monto'] = round(r[k]['monto'], 1)
    return reg


def leer_ppt():
    if not os.path.exists(PPT):
        print('Aviso: no existe "Datos para PPT.xlsx"; obras en ejecución, vivienda y PRESET usarán datos de ejemplo.')
        return {}
    wb = openpyxl.load_workbook(PPT, data_only=True, read_only=True)
    out, avisos = {}, set()
    ws = hoja(wb, 'Obras en ejecucion')
    if ws is not None:
        out['obrasEjecucion'] = leer_obras(ws, avisos)
    ws = hoja(wb, 'Inversiones en vivienda')
    if ws is not None:
        out['vivienda'] = leer_vivienda(ws, avisos)
    ws = hoja(wb, 'Expedientes tecnicos')
    if ws is not None:
        out['preset'] = leer_preset(ws, avisos)
    ws = hoja(wb, 'Obras paralizadas')
    if ws is not None:
        out['obrasParalizadas'] = leer_modalidad(ws, avisos)
    ws = hoja(wb, 'Obras reactivadas')
    if ws is not None:
        out['obrasReactivadas'] = leer_modalidad(ws, avisos)
    if avisos:
        print('Aviso: departamentos no reconocidos (se omitieron):', ', '.join(sorted(avisos)))
    return out


def main():
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(errors='replace')  # tildes seguras en la consola de Windows
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    info = leer_info(wb)
    datos = {
        'fuente': info.get('fuente', 'data/data.xlsx'),
        'anio': info.get('año', ''),
        'descargado': info.get('descargado', ''),
        'generado': datetime.datetime.now().strftime('%Y-%m-%d %H:%M'),
        'ranking': leer_ranking(wb),
        'porRegion': leer_regiones(wb),
    }
    datos.update(leer_ppt())
    js = ('/* Generado desde data/data.xlsx y data/Datos para PPT.xlsx por data/actualizar_data.py. No editar a mano:\n'
          '   actualiza los Excel (o corre data/scrapear_mef.py) y vuelve a correr este script. */\n'
          'window.DATA_REAL = ' + json.dumps(datos, ensure_ascii=False, indent=1) + ';\n')
    tmp = SALIDA + '.tmp'
    io.open(tmp, 'w', encoding='utf-8', newline='\n').write(js)
    os.replace(tmp, SALIDA)
    print('data.js generado:', len(datos['ranking']), 'sectores nacionales,', len(datos['porRegion']), 'regiones del ranking,',
          'obras en ejecución' if 'obrasEjecucion' in datos else 'sin obras en ejecución', '·',
          'vivienda' if 'vivienda' in datos else 'sin vivienda', '·', 'PRESET' if 'preset' in datos else 'sin PRESET', '·',
          'paralizadas' if 'obrasParalizadas' in datos else 'sin paralizadas', '·', 'reactivadas' if 'obrasReactivadas' in datos else 'sin reactivadas')


if __name__ == '__main__':
    main()
