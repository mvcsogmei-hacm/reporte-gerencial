"""
Descarga de la Consulta Amigable del MEF el presupuesto del Gobierno Nacional por sector, separado en
INVERSIONES (proyectos) y ACTIVIDADES, para el total nacional y cada uno de los 25 departamentos,
y lo guarda en data/data.xlsx.

Uso, desde la carpeta dashboard:
    python data/scrapear_mef.py          (año actual)
    python data/scrapear_mef.py 2025     (otro año)
Después:
    python data/actualizar_data.py       (convierte data.xlsx en data/data.js)
y sube los cambios a git.

Hojas que genera en data.xlsx:
  Ranking            Tipo | Sector | PIM | Devengado | Avance %             total del Gobierno Nacional
  Ranking regiones   Tipo | Código | Departamento | Sector | PIM | Devengado | Avance %
  Info               fuente, año, fecha de descarga y direcciones consultadas

"Tipo" es Inversiones (botón Actividades/Proyectos = Proyecto, parámetro ap=Proyecto) o Actividades
(ap=Actividad). Inversiones + actividades = total del MEF por sector (verificado el 11 sep 2026).
"Departamento" es el lugar donde se ubica la meta del gasto (botón Departamento de la
Consulta Amigable), no la oficina que lo ejecuta. La suma de los departamentos da el total nacional.

Necesita openpyxl:  pip install openpyxl
Tarda unos 3 a 5 minutos: hace 52 consultas con pausas, porque el sitio tiene protección contra robots.
"""
import datetime
import html
import os
import re
import sys
import time
import urllib.request

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    sys.exit('Falta openpyxl. Instálalo con: pip install openpyxl')

AQUI = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(AQUI, 'data.xlsx')
BASE = 'https://apps5.mineco.gob.pe/transparencia/Navegador/Navegar_7.aspx'
UA = 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0 Safari/537.36'
PAUSA = 2.5  # segundos entre consultas

DEPARTAMENTOS = [
    ('01', 'Amazonas'), ('02', 'Áncash'), ('03', 'Apurímac'), ('04', 'Arequipa'), ('05', 'Ayacucho'),
    ('06', 'Cajamarca'), ('07', 'Callao'), ('08', 'Cusco'), ('09', 'Huancavelica'), ('10', 'Huánuco'),
    ('11', 'Ica'), ('12', 'Junín'), ('13', 'La Libertad'), ('14', 'Lambayeque'), ('15', 'Lima'),
    ('16', 'Loreto'), ('17', 'Madre de Dios'), ('18', 'Moquegua'), ('19', 'Pasco'), ('20', 'Piura'),
    ('21', 'Puno'), ('22', 'San Martín'), ('23', 'Tacna'), ('24', 'Tumbes'), ('25', 'Ucayali'),
]

# (nombre en data.xlsx, valor del parámetro ap de la Consulta Amigable)
TIPOS = [('Inversiones', 'Proyecto'), ('Actividades', 'Actividad')]


def direccion(anio, depto=None, ap=None):
    """Ruta de la Consulta Amigable: [departamento] > Nivel de gobierno E (Nacional) > Sector, filtrada por
    Actividades/Proyectos (ap=Proyecto: inversiones; ap=Actividad: actividades; sin ap: ambos)."""
    return (BASE + '?_uhc=yes&0=' + ('&21=' + depto if depto else '') +
            '&1=E&2=&y=' + str(anio) + '&cpage=1&psize=400' + ('&ap=' + ap if ap else ''))


def descargar(url, intentos=3):
    for i in range(intentos):
        try:
            req = urllib.request.Request(url, headers={'User-Agent': UA, 'Accept-Language': 'es-PE,es;q=0.9'})
            with urllib.request.urlopen(req, timeout=120) as r:
                return r.read().decode('utf-8', 'replace')
        except Exception as e:  # red caída o bloqueo temporal: reintenta con espera creciente
            if i == intentos - 1:
                raise RuntimeError('No se pudo descargar ' + url + ' · ' + repr(e))
            time.sleep(10 * (i + 1))


def numero(texto):
    t = texto.replace(',', '').strip()
    return float(t) if t not in ('', '-') else 0.0


def leer_cuadro(pagina):
    """Filas del cuadro. Columnas del MEF: nombre, PIA, PIM, Certificación, Compromiso anual,
    Atención de compromiso mensual, Devengado, Girado, Avance %."""
    filas = []
    for tr in re.findall(r'<tr[^>]*>(.*?)</tr>', pagina, re.S):
        if 'name="grp1"' not in tr:
            continue
        celdas = [re.sub(r'\s+', ' ', html.unescape(re.sub(r'<[^>]+>', ' ', c))).strip()
                  for c in re.findall(r'<td[^>]*>(.*?)</td>', tr, re.S)]
        celdas = [c for c in celdas if c]
        if len(celdas) < 9 or ':' not in celdas[0]:
            continue
        pim, dev, av = numero(celdas[2]), numero(celdas[6]), numero(celdas[8])
        if pim > 0 and abs(dev / pim * 100 - av) > 0.2:
            raise RuntimeError('Las columnas del cuadro cambiaron (devengado ÷ PIM no coincide con el avance en "' + celdas[0] + '"). Revisar el script.')
        filas.append({'sector': celdas[0], 'pim': round(pim), 'devengado': round(dev), 'avance': av / 100})
    return filas


def guardar(anio, nacional, regiones, descargado):
    """nacional = {tipo: filas}; regiones = [(código, departamento, {tipo: filas})]."""
    wb = openpyxl.Workbook()
    cab = Font(bold=True, color='FFFFFF')
    fondo = PatternFill('solid', fgColor='1F4E8C')

    def hoja(ws, encabezados, anchos):
        ws.append(encabezados)
        for i, ancho in enumerate(anchos):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = ancho
            ws.cell(row=1, column=i + 1).font = cab
            ws.cell(row=1, column=i + 1).fill = fondo
            ws.cell(row=1, column=i + 1).alignment = Alignment(vertical='center')
        ws.freeze_panes = 'A2'

    ws = wb.active
    ws.title = 'Ranking'
    hoja(ws, ['Tipo', 'Sector', 'PIM', 'Devengado', 'Avance %'], [13, 52, 18, 18, 11])
    for tipo, filas in nacional.items():
        for f in filas:
            ws.append([tipo, f['sector'], f['pim'], f['devengado'], f['avance']])

    ws2 = wb.create_sheet('Ranking regiones')
    hoja(ws2, ['Tipo', 'Código', 'Departamento', 'Sector', 'PIM', 'Devengado', 'Avance %'], [13, 9, 16, 52, 18, 18, 11])
    for tipo, _ in TIPOS:
        for cod, nombre, por_tipo in regiones:
            for f in por_tipo.get(tipo, []):
                ws2.append([tipo, cod, nombre, f['sector'], f['pim'], f['devengado'], f['avance']])

    for w, c_pim, c_av in [(ws, 3, 5), (ws2, 5, 7)]:
        for fila in w.iter_rows(min_row=2):
            fila[c_pim - 1].number_format = '#,##0'
            fila[c_pim].number_format = '#,##0'
            fila[c_av - 1].number_format = '0.0%'

    ws3 = wb.create_sheet('Info')
    hoja(ws3, ['Dato', 'Valor'], [22, 120])
    for fila in [
        ['Fuente', 'MEF · Consulta Amigable (Consulta de Ejecución del Gasto)'],
        ['Nivel de gobierno', 'E: Gobierno Nacional, por sector'],
        ['Tipo', 'Inversiones = Actividades/Proyectos: Proyecto (ap=Proyecto) · Actividades = Actividad (ap=Actividad) · inversiones + actividades = total'],
        ['Año', anio],
        ['Descargado', descargado],
        ['Departamento', 'Lugar donde se ubica la meta del gasto (no la oficina ejecutora)'],
        ['Dirección nacional', direccion(anio, ap='Proyecto') + '   (o ap=Actividad)'],
        ['Dirección por región', direccion(anio, 'XX', 'Proyecto') + '   (XX = código de departamento; o ap=Actividad)'],
    ]:
        ws3.append(fila)

    tmp = XLSX + '.tmp.xlsx'
    wb.save(tmp)
    try:
        os.replace(tmp, XLSX)
    except PermissionError:
        os.remove(tmp)
        sys.exit('No se pudo reemplazar data.xlsx: ciérralo en Excel y vuelve a correr el script.')


def main():
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(errors='replace')  # tildes seguras en la consola de Windows
    anio = int(sys.argv[1]) if len(sys.argv) > 1 else datetime.date.today().year
    print('Consulta Amigable · Gobierno Nacional por sector · inversiones y actividades ·', anio)

    nacional = {}
    for tipo, ap in TIPOS:
        filas = leer_cuadro(descargar(direccion(anio, ap=ap)))
        if len(filas) < 10:
            sys.exit('El cuadro nacional de ' + tipo.lower() + ' trajo ' + str(len(filas)) + ' sectores: la página pudo cambiar o bloquear la consulta. No se modificó data.xlsx.')
        nacional[tipo] = filas
        print('  Nacional · ' + tipo + ':', len(filas), 'sectores')
        time.sleep(PAUSA)

    regiones = []
    for cod, nombre in DEPARTAMENTOS:
        por_tipo = {}
        for tipo, ap in TIPOS:
            time.sleep(PAUSA)
            por_tipo[tipo] = leer_cuadro(descargar(direccion(anio, cod, ap)))
        regiones.append((cod, nombre, por_tipo))
        print('  ' + cod + ' ' + nombre + ':', ' · '.join(t + ' ' + str(len(f)) for t, f in por_tipo.items()), 'sectores')

    # control: en cada tipo, la suma de los departamentos debería dar el total nacional
    for tipo, _ in TIPOS:
        total = {}
        for _, _, por_tipo in regiones:
            for f in por_tipo.get(tipo, []):
                total[f['sector']] = total.get(f['sector'], 0) + f['pim']
        difieren = [f['sector'] for f in nacional[tipo] if f['pim'] and abs(total.get(f['sector'], 0) - f['pim']) / f['pim'] > 0.005]
        if difieren:
            print('  Aviso (' + tipo + '): en', len(difieren), 'sectores la suma de regiones no llega al total nacional (parte del gasto puede estar en el exterior):', ', '.join(difieren))
        else:
            print('  Control (' + tipo + '): la suma de las regiones da el total nacional en todos los sectores.')

    descargado = datetime.datetime.now().strftime('%Y-%m-%d %H:%M')
    guardar(anio, nacional, regiones, descargado)
    print('data.xlsx generado (' + descargado + '). Ahora corre: python data/actualizar_data.py')


if __name__ == '__main__':
    main()
